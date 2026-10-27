#!/usr/bin/env python3
"""Read a payout file, normalize it, run local checks, and list the lookups still needed.

Usage:
  python scripts/prepare.py payouts.csv --outdir ./pdv-work
  python scripts/prepare.py payouts.xlsx --sheet Vendors --map iban="Bank IBAN" --map payee=Vendor
  python scripts/prepare.py payouts.csv --baseline previous.csv --outdir ./pdv-work

Writes <outdir>/prepared.json and prints a short summary. No network calls.
"""
import argparse
import csv
import io
import json
import os
import re
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (  # noqa: E402
    BIC_RE, REASONS, REJECTED, mask_iban, norm_bic, norm_iban, norm_name,
)
from iban_registry import assess_iban, territory_note, typo_hint  # noqa: E402

# Header key -> score. A specific header beats a bare "name" or "account".
ALIAS_SCORES = {
    "payee": {
        "payee": 100, "payeename": 100, "beneficiaryname": 100, "beneficiary": 80,
        "vendorname": 100, "vendor": 80, "suppliername": 100, "supplier": 80,
        "accountholdername": 100, "accountholder": 90, "accountname": 90,
        "recipientname": 100, "recipient": 80, "employeename": 100, "employee": 70,
        "name": 40,
    },
    "iban": {
        "iban": 100, "ibannumber": 100, "ibancode": 100, "bankiban": 100,
        "internationalbankaccountnumber": 100,
    },
    "bic": {
        "bic": 100, "swift": 90, "swiftcode": 100, "biccode": 100, "swiftbic": 100,
        "bicswift": 100, "bankbic": 100, "bankswift": 100,
    },
    "account": {
        "accountnumber": 100, "accountno": 100, "accnumber": 90, "acctnumber": 100,
        "acctno": 90, "bankaccountnumber": 100, "bban": 80, "bankaccount": 50,
        "account": 40,
    },
}


def key_of(header):
    return re.sub(r"[^a-z0-9]", "", (header or "").lower())


def die(message):
    print(message, file=sys.stderr)
    raise SystemExit(1)


def decode_text(path):
    raw = open(path, "rb").read()
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16")
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    for enc in ("utf-8", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    die("Could not decode the file. Save it as UTF-8 CSV and try again.")


def stringify_cell(value):
    """Return (text, was_number). Numbers keep no leading zeros; the caller notes that."""
    if value is None:
        return "", False
    if isinstance(value, bool):
        return str(value), False
    if isinstance(value, int):
        return str(value), True
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value)), True
        return format(value, "f"), True
    return str(value).strip(), False


def find_header(rows):
    """Pick the header from the first 15 non-empty rows. Titles above it are skipped."""
    window = rows[:15]
    if not window:
        die("The file has no data.")
    scores = [header_hits(cells) for _n, cells, _flags in window]
    best = max(scores)
    if best == 0:
        blob = [cell for _n, cells, _flags in window for cell in cells]
        if any(norm_iban(cell)[:2].isalpha() and len(norm_iban(cell)) >= 15 for cell in blob):
            die("This file has no header row. Add a header with columns such as payee, iban and bic.")
        die("Could not find a header row.")
    pos = scores.index(best)
    return pos


def header_hits(cells):
    hits = 0
    seen = set()
    for cell in cells:
        key = key_of(cell)
        if not key or key in seen:
            continue
        if any(key in scores for scores in ALIAS_SCORES.values()):
            hits += 1
            seen.add(key)
    return hits


def read_table(path, sheet=None):
    """Return (headers, data rows, meta). A data row is (source_row, cells, numeric flags)."""
    ext = os.path.splitext(path)[1].lower()
    meta = {"sheet": None}
    if ext in (".xls", ".xlsb"):
        die("Old .xls workbooks are not supported. Save the sheet as .xlsx or CSV.")
    if ext in (".xlsx", ".xlsm"):
        rows, meta["sheet"] = read_excel(path, sheet)
    else:
        rows = read_csv(path)
    if not rows:
        die("The file has no data.")
    header_pos = find_header(rows)
    header_cells = rows[header_pos][1]
    data = rows[header_pos + 1:]
    if not data:
        die("The file has no data rows.")
    width = max([len(header_cells)] + [len(cells) for _n, cells, _f in data])
    headers, seen = [], defaultdict(int)
    for i in range(width):
        raw = header_cells[i] if i < len(header_cells) and header_cells[i] else "column_%d" % (i + 1)
        seen[raw] += 1
        headers.append(raw if seen[raw] == 1 else "%s_%d" % (raw, seen[raw]))
    padded = []
    for n, cells, flags in data:
        cells = (cells + [""] * width)[:width]
        flags = (flags + [False] * width)[:width]
        padded.append((n, cells, flags))
    return headers, padded, meta


def read_excel(path, sheet):
    try:
        import openpyxl
    except ImportError:
        die("Reading .xlsx needs openpyxl (pip install openpyxl). Alternatively export the sheet to CSV.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if sheet:
            if sheet not in names:
                die("No sheet %r. Sheets: %s" % (sheet, ", ".join(names)))
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        rows = []
        for n, record in enumerate(ws.iter_rows(values_only=True), start=1):
            cells, flags = [], []
            for value in record:
                text, was_number = stringify_cell(value)
                cells.append(text)
                flags.append(was_number)
            if any(cells):
                rows.append((n, cells, flags))
        return rows, ws.title
    finally:
        wb.close()


def read_csv(path):
    text = decode_text(path)
    delimiter = None
    start = 1
    first = text.splitlines()[0].strip() if text.strip() else ""
    if re.fullmatch(r"sep=.", first, flags=re.IGNORECASE):
        delimiter = first[-1]
        start = 2
        text = text.split("\n", 1)[1] if "\n" in text else ""
    if delimiter:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    else:
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows = []
    for n, record in enumerate(reader, start=start):
        cells = [c.strip() for c in record]
        if any(cells):
            rows.append((n, cells, [False] * len(cells)))
    return rows


def value_hits(field, header, headers, data_rows):
    idx = headers.index(header)
    hits = 0
    for _n, cells, _flags in data_rows:
        value = cells[idx] if idx < len(cells) else ""
        if field == "iban" and assess_iban(norm_iban(value))["format_ok"]:
            hits += 1
        elif field == "bic" and BIC_RE.match(norm_bic(value)):
            hits += 1
        elif field == "payee" and value.strip():
            hits += 1
        elif field == "account" and value.strip() and not assess_iban(norm_iban(value))["format_ok"]:
            hits += 1
    return hits


def detect_columns(headers, overrides, data_rows):
    """Return (mapping, notes). Overrides win. Ties break toward values that fit the field."""
    mapping = {}
    notes = []
    lowered = {h.lower(): h for h in headers}
    for field, header in overrides.items():
        if header.lower() not in lowered:
            die("--map %s=%s: no such column. Columns are: %s" % (field, header, ", ".join(headers)))
        mapping[field] = lowered[header.lower()]
    for field, scores in ALIAS_SCORES.items():
        if field in mapping:
            continue
        ranked = [(scores[key_of(h)], h) for h in headers if key_of(h) in scores]
        if not ranked:
            continue
        top = max(score for score, _h in ranked)
        tied = [h for score, h in ranked if score == top]
        tied.sort(key=lambda h: (-value_hits(field, h, headers, data_rows), headers.index(h)))
        mapping[field] = tied[0]
    if "iban" not in mapping and "iban" not in overrides:
        claimed = {mapping.get("payee"), mapping.get("bic")}
        best_h, best_n = None, 0
        for h in headers:
            if h in claimed:
                continue
            idx = headers.index(h)
            nonempty = sum(1 for _n, cells, _flags in data_rows if idx < len(cells) and cells[idx].strip())
            n = value_hits("iban", h, headers, data_rows)
            if n < 1 or n * 2 < nonempty:
                continue
            if n > best_n:
                best_h, best_n = h, n
        if best_h:
            if mapping.get("account") == best_h and "account" not in overrides:
                del mapping["account"]
                notes.append("IBAN column taken from values in %r." % best_h)
            elif best_h not in mapping.values():
                notes.append("IBAN column taken from values in %r." % best_h)
            mapping["iban"] = best_h
    return mapping, notes


def finding(code, detail=None):
    return {"code": code, "detail": detail or REASONS[code][1]}


def column_values(headers, table, mapping):
    col = {field: headers.index(header) for field, header in mapping.items()}

    def get(cells, field):
        return cells[col[field]] if field in col else ""

    def was_number(flags, field):
        return bool(field in col and flags[col[field]])

    return get, was_number


def load_baseline(path, sheet, overrides):
    headers, table, _meta = read_table(path, sheet)
    mapping, _notes = detect_columns(headers, overrides, table)
    if "payee" not in mapping or "iban" not in mapping:
        die("The baseline file needs a payee column and an IBAN column. Columns are: %s" % ", ".join(headers))
    get, _was = column_values(headers, table, mapping)
    groups = defaultdict(set)
    for _n, cells, _flags in table:
        key = norm_name(get(cells, "payee"))
        iban = norm_iban(get(cells, "iban"))
        if key and iban:
            groups[key].add(iban)
    return groups


def build_rows(headers, table, mapping):
    get, was_number = column_values(headers, table, mapping)
    rows = []
    for rownum, cells, flags in table:
        payee_raw = get(cells, "payee")
        iban_raw = get(cells, "iban")
        bic_raw = get(cells, "bic")
        acct_raw = get(cells, "account")
        iban, bic = norm_iban(iban_raw), norm_bic(bic_raw)
        findings, notes = [], []

        if was_number(flags, "iban") or was_number(flags, "account"):
            notes.append(
                "A bank cell was stored as a number in Excel, so leading zeros and digits past 15 may already be gone.")

        if not (iban or bic or acct_raw):
            findings.append(finding("MISSING_FIELD"))

        assessed = assess_iban(iban) if iban else None
        if iban and iban_raw.strip() != iban:
            notes.append("Spaces, dots, hyphens or lowercase were removed for the IBAN check. The original value is unchanged.")
        if iban:
            if not assessed["format_ok"]:
                findings.append(finding("IBAN_BAD_FORMAT", assessed["detail"]))
            elif not assessed["checksum_ok"]:
                findings.append(finding("IBAN_BAD_CHECKSUM"))
            elif assessed["qr"]:
                findings.append(finding("QR_IBAN"))
            hint = None if assessed["checksum_ok"] else typo_hint(iban)
            if hint:
                notes.append(
                    "One change would pass the checks (%s). This is a guess, not a correction. Confirm it with the payee."
                    % hint)

        bic_ok = False
        if bic:
            if not BIC_RE.match(bic):
                findings.append(finding("BIC_BAD_FORMAT"))
            else:
                bic_ok = True
            if norm_bic(bic_raw) == bic and bic_raw.strip() != bic:
                notes.append("Spaces or hyphens were removed for the BIC check. The original value is unchanged.")

        iban_ok = bool(assessed and assessed["checksum_ok"])
        if iban_ok and bic_ok and bic[4:6] != iban[:2]:
            note = territory_note(iban[:2], bic[4:6])
            if note:
                notes.append(note)
            else:
                findings.append(finding(
                    "COUNTRY_MISMATCH",
                    "BIC country is %s but IBAN country is %s." % (bic[4:6], iban[:2])))

        rows.append({
            "row": rownum,
            "original": dict(zip(headers, cells)),
            "payee": payee_raw,
            "iban": iban,
            "bic": bic,
            "account": acct_raw,
            "iban_format_ok": bool(assessed and assessed["format_ok"]),
            "iban_ok_local": iban_ok,
            "bic_ok_local": bic_ok,
            "findings": findings,
            "notes": notes,
        })
    return rows


def add_duplicates(rows):
    def acct_key(row):
        if row["iban"]:
            return row["iban"]
        if row["account"]:
            return "ACCT:" + row["account"].replace(" ", "").upper()
        return ""

    by_pair = defaultdict(list)
    by_iban = defaultdict(set)
    by_payee = defaultdict(set)
    for i, row in enumerate(rows):
        key = acct_key(row)
        if not key:
            continue
        by_pair[(norm_name(row["payee"]), key)].append(i)
        if row["iban"]:
            by_iban[row["iban"]].add(norm_name(row["payee"]))
        if norm_name(row["payee"]) and row["iban"]:
            by_payee[norm_name(row["payee"])].add(row["iban"])

    for (_name, _key), idxs in by_pair.items():
        if len(idxs) < 2:
            continue
        lines = ", ".join(str(rows[i]["row"]) for i in idxs)
        for i in idxs:
            rows[i]["findings"].append(finding(
                "DUPLICATE_ROW",
                "Same payee and account on rows %s. Paying all of them would pay twice." % lines))
    for row in rows:
        if row["iban"] and len(by_iban[row["iban"]]) > 1:
            row["findings"].append(finding(
                "DUPLICATE_IBAN",
                "This IBAN appears under %d different payee names." % len(by_iban[row["iban"]])))
        key = norm_name(row["payee"])
        if key and len(by_payee.get(key, ())) > 1:
            row["findings"].append(finding(
                "MULTIPLE_IBANS_PER_PAYEE",
                "This payee appears with %d different IBANs." % len(by_payee[key])))


def add_baseline(rows, baseline):
    """Flag a payee whose current IBAN is not one of the baseline IBANs for that exact name."""
    changes = 0
    for row in rows:
        key = norm_name(row["payee"])
        if not key or not row["iban"] or key not in baseline:
            continue
        if row["iban"] in baseline[key]:
            continue
        row["findings"].append(finding("IBAN_CHANGED"))
        changes += 1
    return changes


def locally_rejected(row):
    return any(REASONS[item["code"]][0] == REJECTED for item in row["findings"])


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--outdir", default="pdv-work")
    ap.add_argument("--sheet")
    ap.add_argument("--baseline", help="Previous payout file. Same payee name, different IBAN, is IBAN_CHANGED.")
    ap.add_argument("--baseline-sheet")
    ap.add_argument("--map", action="append", default=[], metavar="FIELD=COLUMN",
                    help="Override column detection. FIELD is payee, iban, bic or account.")
    args = ap.parse_args(argv)

    overrides = {}
    for item in args.map:
        field, _, header = item.partition("=")
        if field not in ALIAS_SCORES or not header:
            die("--map must look like iban=\"Column name\" (fields: payee, iban, bic, account)")
        overrides[field] = header

    headers, table, meta = read_table(args.input, args.sheet)
    mapping, detection_notes = detect_columns(headers, overrides, table)
    if not any(field in mapping for field in ("iban", "bic", "account")):
        die("Could not find an IBAN, BIC or account column. Columns are: %s\n"
            "Re-run with --map iban=\"<column>\" (and bic/payee/account as needed)." % ", ".join(headers))

    rows = build_rows(headers, table, mapping)
    add_duplicates(rows)
    baseline_changes = 0
    if args.baseline:
        baseline = load_baseline(args.baseline, args.baseline_sheet, overrides)
        baseline_changes = add_baseline(rows, baseline)

    ibans, bics = [], []
    for row in rows:
        if locally_rejected(row):
            continue
        if row["iban_ok_local"] and row["iban"] not in ibans:
            ibans.append(row["iban"])
        if row["bic_ok_local"] and row["bic"] not in bics:
            bics.append(row["bic"])

    os.makedirs(args.outdir, exist_ok=True)
    out = {
        "source": os.path.abspath(args.input),
        "headers": headers,
        "columns": mapping,
        "rows": rows,
        "to_lookup": {"ibans": ibans, "bics": bics},
    }
    dest = os.path.join(args.outdir, "prepared.json")
    with open(dest, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)

    local_rejects = sum(1 for row in rows if locally_rejected(row))
    print("Rows read: %d" % len(rows))
    if meta.get("sheet"):
        print("Sheet: " + meta["sheet"])
    print("Columns used: " + ", ".join("%s -> %s" % (field, header) for field, header in mapping.items()))
    for note in detection_notes:
        print(note)
    missing = [field for field in ALIAS_SCORES if field not in mapping]
    if missing:
        print("Columns not found: " + ", ".join(missing))
    print("Rejected locally (no lookup needed): %d" % local_rejects)
    if args.baseline:
        print("Baseline IBAN changes: %d" % baseline_changes)
    print("Lookups needed: %d unique IBANs, %d unique BICs" % (len(ibans), len(bics)))
    if ibans:
        print("Example IBAN (masked): " + mask_iban(ibans[0]))
    print("Wrote " + dest)
    return 0


if __name__ == "__main__":
    main()
